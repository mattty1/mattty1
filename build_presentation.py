#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generator kolorowej prezentacji PowerPoint (styl Canva).

Temat: Wizerunek marek iPhone vs Samsung w opinii konsumentów.
Na podstawie raportu z badania marketingowego (UMCS, Wydział Ekonomiczny).

Uruchomienie:
    python3 build_presentation.py
Wynik:
    Prezentacja_iPhone_vs_Samsung.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn

# ---------------------------------------------------------------------------
# PALETA KOLORÓW (barwy marek Apple / Samsung)
#   Apple  -> grafit / czerń / srebro + jasnoniebieski akcent
#   Samsung-> głęboki granatowy niebieski + jasne niebieskie tony
# Nazwy zmiennych pozostają jak wcześniej (mapowanie ról), aby nie zmieniać
# całego układu — zmieniają się jedynie wartości kolorów.
# ---------------------------------------------------------------------------
INK       = RGBColor(0x1D, 0x1D, 0x1F)   # Apple "space black" (tekst)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CLOUD     = RGBColor(0xF5, 0xF5, 0xF7)   # Apple srebrne tło
PURPLE    = RGBColor(0x00, 0x71, 0xE3)   # Apple blue — główny akcent
PURPLE_DK = RGBColor(0x0A, 0x12, 0x2A)   # ciemny granat (tła ciemnych slajdów)
PINK      = RGBColor(0x14, 0x28, 0xA0)   # Samsung blue — akcent drugi
CORAL     = RGBColor(0x48, 0x48, 0x4A)   # Apple grafit
AMBER     = RGBColor(0x5A, 0xC8, 0xFA)   # Apple jasnoniebieski (wyróżnienie)
TEAL      = RGBColor(0x00, 0x9D, 0xDC)   # cyjanowy niebieski
SKY       = RGBColor(0x1B, 0x6E, 0xF3)   # jasny Samsung blue
GREEN     = RGBColor(0x0A, 0x84, 0xFF)   # żywy niebieski (status: potwierdzona)
GREY      = RGBColor(0x86, 0x86, 0x8B)   # Apple szary (tekst drugorzędny)

IPHONE_C  = RGBColor(0x1D, 0x1D, 0x1F)   # grafit Apple
SAMSUNG_C = RGBColor(0x14, 0x28, 0xA0)   # niebieski Samsung

ACCENTS = [PURPLE, PINK, TEAL, AMBER, CORAL, SKY]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------------------
# HELPERY
# ---------------------------------------------------------------------------
def _gradient(shape, c1, c2, angle=45):
    """Gradient liniowy na kształcie (dwa stopnie)."""
    shape.fill.solid()
    sp = shape.fill._xPr
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    grad = sp.makeelement(qn('a:gradFill'), {})
    gsLst = grad.makeelement(qn('a:gsLst'), {})

    def gs(pos, color):
        g = grad.makeelement(qn('a:gs'), {'pos': str(pos)})
        clr = grad.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (color[0], color[1], color[2])})
        g.append(clr)
        return g

    gsLst.append(gs(0, c1))
    gsLst.append(gs(100000, c2))
    grad.append(gsLst)
    lin = grad.makeelement(qn('a:lin'), {'ang': str(int(angle * 60000)), 'scaled': '1'})
    grad.append(lin)
    ln = sp.find(qn('a:ln'))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp.append(grad)
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, color=None, shape=MSO_SHAPE.RECTANGLE,
             grad=None, grad_angle=45, line=None, line_w=0, shadow=False):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    if grad is not None:
        _gradient(sh, grad[0], grad[1], grad_angle)
    elif color is not None:
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
    else:
        sh.fill.background()
    if line is not None:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w if line_w else 1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if shadow:
        _soft_shadow(sh)
    return sh


def _soft_shadow(shape):
    spPr = shape._element.spPr
    effectLst = spPr.makeelement(qn('a:effectLst'), {})
    outer = spPr.makeelement(qn('a:outerShdw'), {
        'blurRad': '90000', 'dist': '40000', 'dir': '5400000', 'rotWithShape': '0'})
    clr = spPr.makeelement(qn('a:srgbClr'), {'val': '2A1A55'})
    alpha = spPr.makeelement(qn('a:alpha'), {'val': '26000'})
    clr.append(alpha)
    outer.append(clr)
    effectLst.append(outer)
    spPr.append(effectLst)


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=6, line_spacing=1.0, wrap=True):
    """runs: lista akapitów; każdy akapit to lista (tekst, dict_opcji)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        if isinstance(para, tuple):
            para = [para]
        for text, opt in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = opt.get('font', 'Calibri')
            f.size = Pt(opt.get('size', 18))
            f.bold = opt.get('bold', False)
            f.italic = opt.get('italic', False)
            f.color.rgb = opt.get('color', INK)
    return tb


def add_circle(slide, x, y, d, color=None, grad=None, grad_angle=45):
    return add_rect(slide, x, y, d, d, color=color, shape=MSO_SHAPE.OVAL, grad=grad, grad_angle=grad_angle)


def page_number(slide, n, dark=False):
    add_text(slide, SW - Inches(0.9), SH - Inches(0.55), Inches(0.6), Inches(0.35),
             [[(f"{n:02d}", {'size': 12, 'bold': True, 'color': WHITE if dark else GREY})]],
             align=PP_ALIGN.RIGHT)


def kicker(slide, text, color, x=Inches(0.9), y=Inches(0.62)):
    """Mały kolorowy 'badge' nad tytułem sekcji."""
    w = Inches(0.22 + 0.108 * len(text))
    pill = add_rect(slide, x, y, w, Inches(0.42), color=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    pill.adjustments[0] = 0.5
    add_text(slide, x, y, w, Inches(0.42),
             [[(text.upper(), {'size': 12.5, 'bold': True, 'color': WHITE, 'font': 'Calibri'})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return y + Inches(0.42)


def section_title(slide, title, color, y=Inches(1.18)):
    add_text(slide, Inches(0.9), y, Inches(11.5), Inches(1.0),
             [[(title, {'size': 34, 'bold': True, 'color': INK})]])
    add_rect(slide, Inches(0.92), y + Inches(0.92), Inches(0.9), Inches(0.09),
             color=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def base_slide(bg=CLOUD):
    slide = prs.slides.add_slide(BLANK)
    add_rect(slide, 0, 0, SW, SH, color=bg)
    return slide


def style_chart_text(chart, size=11, color=INK):
    chart.font.size = Pt(size)
    chart.font.color.rgb = color
    chart.font.name = 'Calibri'


# ===========================================================================
# SLAJD 1 — STRONA TYTUŁOWA
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, grad=(PURPLE_DK, PURPLE), grad_angle=60)
add_circle(s, Inches(-1.6), Inches(-1.8), Inches(4.2), grad=(PINK, CORAL))
add_circle(s, SW - Inches(2.4), SH - Inches(2.6), Inches(4.6), grad=(TEAL, SKY))
add_circle(s, SW - Inches(1.0), Inches(0.6), Inches(1.1), color=AMBER)
add_circle(s, Inches(1.2), SH - Inches(1.4), Inches(0.8), color=PINK)

add_text(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(0.9),
         [[("UNIWERSYTET MARII CURIE-SKŁODOWSKIEJ W LUBLINIE", {'size': 14, 'bold': True, 'color': WHITE})],
          [("Wydział Ekonomiczny  |  Kierunek: Zarządzanie", {'size': 12.5, 'color': RGBColor(0xD9, 0xCB, 0xF7)})]],
         space_after=2)

add_circle(s, Inches(0.9), Inches(2.35), Inches(0.95), color=IPHONE_C)
add_text(s, Inches(0.9), Inches(2.35), Inches(0.95), Inches(0.95),
         [[("iOS", {'size': 17, 'bold': True, 'color': WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_circle(s, Inches(2.0), Inches(2.35), Inches(0.95), color=SAMSUNG_C)
add_text(s, Inches(2.0), Inches(2.35), Inches(0.95), Inches(0.95),
         [[("SAM", {'size': 14, 'bold': True, 'color': WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

add_text(s, Inches(0.9), Inches(3.45), Inches(11.6), Inches(2.0),
         [[("iPhone ", {'size': 52, 'bold': True, 'color': WHITE}),
           ("vs", {'size': 40, 'bold': True, 'italic': True, 'color': AMBER}),
           (" Samsung", {'size': 52, 'bold': True, 'color': WHITE})],
          [("Wizerunek marek a decyzje zakupowe konsumentów", {'size': 22, 'color': RGBColor(0xE7, 0xDE, 0xFB)})]],
         space_after=8)

add_rect(s, Inches(0.9), Inches(6.05), Inches(11.55), Inches(0.95),
         color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.25
add_text(s, Inches(1.2), Inches(6.05), Inches(11.0), Inches(0.95),
         [[("AUTORKI:  ", {'size': 13, 'bold': True, 'color': PURPLE}),
           ("Magdalena Gałek  •  Aleksandra Macioszek  •  Łucja Marcyniuk  •  Julia Nasiłowska  •  Dominika Zdun",
            {'size': 13, 'color': INK})],
          [("Raport z badania marketingowego  •  maj 2026", {'size': 11.5, 'italic': True, 'color': GREY})]],
         anchor=MSO_ANCHOR.MIDDLE, space_after=2)


# ===========================================================================
# SLAJD 2 — AGENDA + CEL BADANIA
# ===========================================================================
s = base_slide()
kicker(s, "Wprowadzenie", PURPLE)
section_title(s, "O czym jest to badanie?", PURPLE)

add_rect(s, Inches(0.9), Inches(2.45), Inches(6.0), Inches(4.3),
         color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True).adjustments[0] = 0.06
add_text(s, Inches(1.25), Inches(2.75), Inches(5.3), Inches(3.8),
         [[("Cel badania", {'size': 20, 'bold': True, 'color': PINK})],
          [("Poznanie sposobu postrzegania marek iPhone (Apple) oraz Samsung przez konsumentów oraz określenie wpływu ich wizerunku na deklaracje zakupowe.",
            {'size': 16, 'color': INK})],
          [("", {'size': 6})],
          [("Główne pytanie badawcze:", {'size': 14, 'bold': True, 'color': PURPLE})],
          [("Jakie są opinie konsumentów na temat wizerunku marek iPhone i Samsung oraz jak wpływają one na ich deklaracje zakupowe?",
            {'size': 14, 'italic': True, 'color': GREY})]],
         space_after=8, line_spacing=1.05)

plan = [
    ("01", "Metodyka badania", PURPLE),
    ("02", "Hipotezy badawcze", PINK),
    ("03", "Próba badawcza (N=81)", TEAL),
    ("04", "Wyniki badania", AMBER),
    ("05", "Wnioski i zalecenia", CORAL),
]
py = Inches(2.45)
for num, txt, col in plan:
    add_rect(s, Inches(7.2), py, Inches(5.2), Inches(0.7),
             color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True).adjustments[0] = 0.3
    add_circle(s, Inches(7.32), py + Inches(0.1), Inches(0.5), color=col)
    add_text(s, Inches(7.32), py + Inches(0.1), Inches(0.5), Inches(0.5),
             [[(num, {'size': 13, 'bold': True, 'color': WHITE})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(8.0), py, Inches(4.3), Inches(0.7),
             [[(txt, {'size': 16, 'bold': True, 'color': INK})]],
             anchor=MSO_ANCHOR.MIDDLE)
    py += Inches(0.85)
page_number(s, 2)


# ===========================================================================
# SLAJD 3 — METODYKA: PROBLEM I PYTANIA SZCZEGÓŁOWE
# ===========================================================================
s = base_slide()
kicker(s, "Rozdział 1 — Metodyka", TEAL)
section_title(s, "Pytania szczegółowe badania", TEAL)

questions = [
    ("PS1", "Które cechy (aparat, prestiż, design, bateria, cena) są najsilniejszymi atrybutami iPhone'a, a które Samsunga?", PURPLE),
    ("PS2", "Czy iPhone jest postrzegany jako symbol wysokiego statusu społecznego?", PINK),
    ("PS3", "W jakim stopniu zmiana systemu (Android vs iOS) wpływa na lojalność wobec marki?", TEAL),
    ("PS4", "Jakie decyzje zakupowe podjęliby konsumenci przy braku ograniczeń budżetowych?", AMBER),
]
gx, gy = Inches(0.9), Inches(2.55)
cw, ch = Inches(5.7), Inches(1.9)
for i, (tag, txt, col) in enumerate(questions):
    x = gx + (cw + Inches(0.15)) * (i % 2)
    y = gy + (ch + Inches(0.2)) * (i // 2)
    card = add_rect(s, x, y, cw, ch, color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    card.adjustments[0] = 0.07
    add_rect(s, x, y, Inches(0.16), ch, color=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x + Inches(0.4), y + Inches(0.2), Inches(2.0), Inches(0.6),
             [[(tag, {'size': 22, 'bold': True, 'color': col})]])
    add_text(s, x + Inches(0.4), y + Inches(0.78), cw - Inches(0.7), Inches(1.0),
             [[(txt, {'size': 14.5, 'color': INK})]], line_spacing=1.05)
page_number(s, 3)


# ===========================================================================
# SLAJD 4 — HIPOTEZY BADAWCZE
# ===========================================================================
s = base_slide(bg=PURPLE_DK)
add_circle(s, SW - Inches(2.0), SH - Inches(2.2), Inches(3.6), grad=(PINK, CORAL))
add_circle(s, Inches(-1.0), Inches(-1.2), Inches(2.6), grad=(TEAL, SKY))
kicker(s, "Rozdział 1 — Metodyka", AMBER)
add_text(s, Inches(0.9), Inches(1.18), Inches(11.5), Inches(1.0),
         [[("Hipotezy badawcze", {'size': 34, 'bold': True, 'color': WHITE})]])
add_rect(s, Inches(0.92), Inches(2.1), Inches(0.9), Inches(0.09), color=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

hyps = [
    ("H1", "iPhone wyżej w prestiżu, jakości i designie; Samsung lepszy w cenie i baterii.", PINK),
    ("H2", "Większość nie zgadza się, że iPhone to symbol wysokiego statusu społecznego.", TEAL),
    ("H3", "Zmiana systemu operacyjnego jest dla konsumentów mało istotna.", AMBER),
    ("H4", "Przy nieograniczonym budżecie konsumenci częściej wybierają iPhone'a.", CORAL),
]
hy = Inches(2.6)
for tag, txt, col in hyps:
    add_rect(s, Inches(0.9), hy, Inches(11.5), Inches(0.92),
             color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.18
    add_circle(s, Inches(1.05), hy + Inches(0.13), Inches(0.66), color=col)
    add_text(s, Inches(1.05), hy + Inches(0.13), Inches(0.66), Inches(0.66),
             [[(tag, {'size': 17, 'bold': True, 'color': WHITE})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.95), hy, Inches(10.3), Inches(0.92),
             [[(txt, {'size': 16, 'color': INK})]], anchor=MSO_ANCHOR.MIDDLE)
    hy += Inches(1.05)
page_number(s, 4, dark=True)


# ===========================================================================
# SLAJD 5 — METODY BADAWCZE (CAWI, KWESTIONARIUSZ)
# ===========================================================================
s = base_slide()
kicker(s, "Rozdział 1 — Metodyka", PINK)
section_title(s, "Jak prowadziłyśmy badanie?", PINK)

stats = [
    ("CAWI", "Technika", "Wspomagany komputerowo wywiad internetowy", PURPLE),
    ("17", "Pytań", "13 merytorycznych + 4 metryczkowe", TEAL),
    ("maj 2026", "Termin", "Sondaż diagnostyczny, dobór celowo-wygodny", AMBER),
]
sx = Inches(0.9)
for big, mid, small, col in stats:
    card = add_rect(s, sx, Inches(2.5), Inches(3.7), Inches(2.3),
                    grad=(col, RGBColor(min(col[0]+30, 255), min(col[1]+30, 255), min(col[2]+40, 255))),
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    card.adjustments[0] = 0.08
    add_text(s, sx + Inches(0.3), Inches(2.7), Inches(3.1), Inches(1.0),
             [[(big, {'size': 38, 'bold': True, 'color': WHITE})]])
    add_text(s, sx + Inches(0.3), Inches(3.65), Inches(3.1), Inches(0.5),
             [[(mid, {'size': 17, 'bold': True, 'color': WHITE})]])
    add_text(s, sx + Inches(0.3), Inches(4.05), Inches(3.1), Inches(0.7),
             [[(small, {'size': 12.5, 'color': RGBColor(0xF3, 0xEE, 0xFF)})]], line_spacing=1.0)
    sx += Inches(3.95)

add_text(s, Inches(0.9), Inches(5.15), Inches(11.5), Inches(0.5),
         [[("Rodzaje zastosowanych pytań:", {'size': 16, 'bold': True, 'color': INK})]])
types = ["Filtrujące / dychotomiczne", "Otwarte (asocjacyjne)", "Skale Likerta 1-5", "Dylematowe (intencje zakupu)"]
tx = Inches(0.9)
for i, t in enumerate(types):
    col = ACCENTS[i % len(ACCENTS)]
    w = Inches(0.5 + 0.105 * len(t))
    pill = add_rect(s, tx, Inches(5.65), w, Inches(0.55), color=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    pill.adjustments[0] = 0.5
    add_text(s, tx, Inches(5.65), w, Inches(0.55),
             [[(t, {'size': 13, 'bold': True, 'color': WHITE})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tx += w + Inches(0.2)
page_number(s, 5)


# ===========================================================================
# SLAJD 6 — CHARAKTERYSTYKA PRÓBY (N=81) + WYKRESY
# ===========================================================================
s = base_slide()
kicker(s, "Rozdział 1 — Próba", SKY)
section_title(s, "Kim byli respondenci? (N=81)", SKY)

badge = add_rect(s, Inches(10.6), Inches(0.55), Inches(1.9), Inches(1.2),
                 grad=(PINK, CORAL), shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
badge.adjustments[0] = 0.18
add_text(s, Inches(10.6), Inches(0.62), Inches(1.9), Inches(1.1),
         [[("81", {'size': 40, 'bold': True, 'color': WHITE})],
          [("respondentów", {'size': 12, 'color': WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

cd = CategoryChartData()
cd.categories = ['Kobiety', 'Mężczyźni', 'Inna / brak odp.']
cd.add_series('Płeć', (61.73, 32.10, 6.17))
gframe = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.9), Inches(2.5),
                            Inches(5.6), Inches(4.4), cd)
ch1 = gframe.chart
ch1.has_title = True
ch1.chart_title.text_frame.text = "Struktura wg płci (%)"
ch1.chart_title.text_frame.paragraphs[0].font.size = Pt(15)
ch1.chart_title.text_frame.paragraphs[0].font.bold = True
ch1.chart_title.text_frame.paragraphs[0].font.color.rgb = INK
ch1.has_legend = True
ch1.legend.position = XL_LEGEND_POSITION.BOTTOM
ch1.legend.include_in_layout = False
plot = ch1.plots[0]
plot.has_data_labels = True
plot.data_labels.number_format = '0.0"%"'
plot.data_labels.number_format_is_linked = False
plot.data_labels.font.size = Pt(11)
plot.data_labels.font.bold = True
plot.data_labels.font.color.rgb = WHITE
for i, c in enumerate([PINK, SKY, AMBER]):
    plot.series[0].points[i].format.fill.solid()
    plot.series[0].points[i].format.fill.fore_color.rgb = c
style_chart_text(ch1)

cd2 = CategoryChartData()
cd2.categories = ['<18', '18-24', '25-34', '35-44', '45-54', '55+']
cd2.add_series('Wiek', (11.11, 67.90, 2.47, 7.41, 8.64, 2.47))
g2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.8), Inches(2.5),
                        Inches(5.7), Inches(4.4), cd2)
ch2 = g2.chart
ch2.has_title = True
ch2.chart_title.text_frame.text = "Struktura wg wieku (%)"
ch2.chart_title.text_frame.paragraphs[0].font.size = Pt(15)
ch2.chart_title.text_frame.paragraphs[0].font.bold = True
ch2.chart_title.text_frame.paragraphs[0].font.color.rgb = INK
ch2.has_legend = False
plot2 = ch2.plots[0]
plot2.has_data_labels = True
plot2.data_labels.number_format = '0.0"%"'
plot2.data_labels.number_format_is_linked = False
plot2.data_labels.font.size = Pt(10)
plot2.data_labels.font.bold = True
plot2.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot2.gap_width = 60
ser2 = plot2.series[0]
for i in range(6):
    ser2.points[i].format.fill.solid()
    ser2.points[i].format.fill.fore_color.rgb = ACCENTS[i % len(ACCENTS)]
ch2.value_axis.has_major_gridlines = False
ch2.value_axis.visible = False
style_chart_text(ch2, size=11)

add_text(s, Inches(0.9), Inches(6.95), Inches(11.5), Inches(0.4),
         [[("Dominują młodzi (18-24 lata: 67,9%), uczniowie/studenci (49,4%) — najaktywniejsi użytkownicy nowych technologii.",
            {'size': 12.5, 'italic': True, 'color': GREY})]])
page_number(s, 6)


# ===========================================================================
# SLAJD 7 — WYNIKI 2.1: OCENA WIZERUNKU MAREK
# ===========================================================================
s = base_slide()
kicker(s, "Rozdział 2 — Wyniki", PURPLE)
section_title(s, "Ocena wizerunku marek (skala 1-5)", PURPLE)

cd3 = CategoryChartData()
cd3.categories = ['Aparat', 'Prestiż', 'Jakość/design', 'Niezawodność',
                  'Intuicyjność', 'Cena/jakość', 'Innowacyjność', 'Bateria']
cd3.add_series('iPhone', (4.51, 4.27, 4.20, 3.56, 3.67, 3.17, 3.94, 2.80))
cd3.add_series('Samsung', (3.78, 3.59, 3.83, 3.77, 3.88, 3.93, 3.65, 4.00))
g3 = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.9), Inches(2.4),
                        Inches(8.3), Inches(4.7), cd3)
ch3 = g3.chart
ch3.has_title = False
ch3.has_legend = True
ch3.legend.position = XL_LEGEND_POSITION.TOP
ch3.legend.include_in_layout = False
plot3 = ch3.plots[0]
plot3.gap_width = 60
plot3.has_data_labels = True
plot3.data_labels.number_format = '0.00'
plot3.data_labels.number_format_is_linked = False
plot3.data_labels.font.size = Pt(9)
plot3.data_labels.font.bold = True
plot3.series[0].format.fill.solid()
plot3.series[0].format.fill.fore_color.rgb = IPHONE_C
plot3.series[1].format.fill.solid()
plot3.series[1].format.fill.fore_color.rgb = SAMSUNG_C
ch3.value_axis.minimum_scale = 0
ch3.value_axis.maximum_scale = 5
ch3.value_axis.has_major_gridlines = True
style_chart_text(ch3, size=11)

add_rect(s, Inches(9.45), Inches(2.4), Inches(3.0), Inches(4.7),
         color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True).adjustments[0] = 0.06
add_text(s, Inches(9.75), Inches(2.65), Inches(2.5), Inches(4.3),
         [[("H1 potwierdzona", {'size': 16, 'bold': True, 'color': GREEN})],
          [("", {'size': 4})],
          [("iPhone wygrywa:", {'size': 13, 'bold': True, 'color': IPHONE_C})],
          [("aparat (4,51), prestiż (4,27), jakość/design (4,20)", {'size': 12.5, 'color': INK})],
          [("", {'size': 4})],
          [("Samsung wygrywa:", {'size': 13, 'bold': True, 'color': SAMSUNG_C})],
          [("bateria (4,00), cena/jakość (3,93), intuicyjność (3,88)", {'size': 12.5, 'color': INK})],
          [("", {'size': 4})],
          [("iPhone = premium, Samsung = funkcjonalność i ekonomia.", {'size': 12.5, 'italic': True, 'color': GREY})]],
         space_after=4, line_spacing=1.03)
page_number(s, 7)


# ===========================================================================
# SLAJD 8 — WYNIKI 2.2: STATUS SPOŁECZNY I BARIERA TECHNOLOGICZNA
# ===========================================================================
s = base_slide()
kicker(s, "Rozdział 2 — Wyniki", PINK)
section_title(s, "Status społeczny i bariera technologiczna", PINK)

cd4 = CategoryChartData()
cd4.categories = ['Nie zgadza się', 'Trudno powiedzieć', 'Zgadza się']
cd4.add_series('Status', (50.6, 22.2, 27.2))
g4 = s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(0.9), Inches(2.6),
                        Inches(5.4), Inches(4.0), cd4)
ch4 = g4.chart
ch4.has_title = True
ch4.chart_title.text_frame.text = "Czy iPhone = wysoki status społeczny? (%)"
ch4.chart_title.text_frame.paragraphs[0].font.size = Pt(13.5)
ch4.chart_title.text_frame.paragraphs[0].font.bold = True
ch4.chart_title.text_frame.paragraphs[0].font.color.rgb = INK
ch4.has_legend = True
ch4.legend.position = XL_LEGEND_POSITION.BOTTOM
ch4.legend.include_in_layout = False
p4 = ch4.plots[0]
p4.has_data_labels = True
p4.data_labels.number_format = '0.0"%"'
p4.data_labels.number_format_is_linked = False
p4.data_labels.font.size = Pt(12)
p4.data_labels.font.bold = True
p4.data_labels.font.color.rgb = WHITE
for i, c in enumerate([GREEN, GREY, PINK]):
    p4.series[0].points[i].format.fill.solid()
    p4.series[0].points[i].format.fill.fore_color.rgb = c
style_chart_text(ch4)

add_rect(s, Inches(6.7), Inches(2.6), Inches(5.75), Inches(2.0),
         color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True).adjustments[0] = 0.08
add_text(s, Inches(7.0), Inches(2.8), Inches(5.2), Inches(1.7),
         [[("Bariera technologiczna (zmiana Android <-> iOS)", {'size': 15, 'bold': True, 'color': PURPLE})],
          [("37,0% — to nie jest duży problem", {'size': 14, 'color': INK})],
          [("34,6% — to istotne utrudnienie", {'size': 14, 'color': INK})],
          [("28,4% — trudno powiedzieć", {'size': 14, 'color': GREY})]],
         space_after=4, line_spacing=1.05)

add_rect(s, Inches(6.7), Inches(4.8), Inches(2.8), Inches(1.8),
         grad=(GREEN, TEAL), shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True).adjustments[0] = 0.1
add_text(s, Inches(6.9), Inches(5.0), Inches(2.4), Inches(1.5),
         [[("H2", {'size': 24, 'bold': True, 'color': WHITE})],
          [("POTWIERDZONA", {'size': 13, 'bold': True, 'color': WHITE})],
          [("50,6% nie łączy iPhone'a ze statusem", {'size': 11.5, 'color': WHITE})]],
         space_after=3, line_spacing=1.0)
add_rect(s, Inches(9.65), Inches(4.8), Inches(2.8), Inches(1.8),
         grad=(AMBER, CORAL), shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True).adjustments[0] = 0.1
add_text(s, Inches(9.85), Inches(5.0), Inches(2.4), Inches(1.5),
         [[("H3", {'size': 24, 'bold': True, 'color': WHITE})],
          [("CZĘŚCIOWO", {'size': 13, 'bold': True, 'color': WHITE})],
          [("duże zróżnicowanie odpowiedzi", {'size': 11.5, 'color': WHITE})]],
         space_after=3, line_spacing=1.0)
page_number(s, 8)


# ===========================================================================
# SLAJD 9 — WYNIKI 2.3: PREFERENCJE ZAKUPOWE
# ===========================================================================
s = base_slide(bg=PURPLE_DK)
add_circle(s, Inches(-1.2), SH - Inches(1.8), Inches(3.2), grad=(TEAL, SKY))
add_circle(s, SW - Inches(1.6), Inches(-0.8), Inches(2.4), grad=(PINK, CORAL))
kicker(s, "Rozdział 2 — Wyniki", AMBER)
add_text(s, Inches(0.9), Inches(1.18), Inches(11.5), Inches(1.0),
         [[("Preferencje zakupowe respondentów", {'size': 34, 'bold': True, 'color': WHITE})]])
add_rect(s, Inches(0.92), Inches(2.1), Inches(0.9), Inches(0.09), color=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

add_text(s, Inches(0.9), Inches(2.55), Inches(4.4), Inches(2.6),
         [[("63%", {'size': 100, 'bold': True, 'color': AMBER})]],
         anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(0.95), Inches(4.85), Inches(4.4), Inches(1.6),
         [[("respondentów wybiera iPhone'a", {'size': 18, 'bold': True, 'color': WHITE})],
          [("przy nieograniczonym budżecie", {'size': 14, 'color': RGBColor(0xD9, 0xCB, 0xF7)})]],
         space_after=2)

cd5 = CategoryChartData()
cd5.categories = ['iPhone', 'Samsung', 'Inna marka']
cd5.add_series('Wybór', (63.0, 27.2, 9.8))
g5 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.6), Inches(2.5),
                        Inches(6.9), Inches(4.4), cd5)
ch5 = g5.chart
ch5.has_title = True
ch5.chart_title.text_frame.text = "Wybór smartfona przy nieograniczonym budżecie (%)"
ch5.chart_title.text_frame.paragraphs[0].font.size = Pt(13)
ch5.chart_title.text_frame.paragraphs[0].font.bold = True
ch5.chart_title.text_frame.paragraphs[0].font.color.rgb = WHITE
ch5.has_legend = False
p5 = ch5.plots[0]
p5.gap_width = 80
p5.has_data_labels = True
p5.data_labels.number_format = '0.0"%"'
p5.data_labels.number_format_is_linked = False
p5.data_labels.font.size = Pt(14)
p5.data_labels.font.bold = True
p5.data_labels.font.color.rgb = WHITE
p5.data_labels.position = XL_LABEL_POSITION.INSIDE_END
for i, c in enumerate([AMBER, SKY, PINK]):
    p5.series[0].points[i].format.fill.solid()
    p5.series[0].points[i].format.fill.fore_color.rgb = c
ch5.value_axis.visible = False
ch5.value_axis.has_major_gridlines = False
ch5.category_axis.tick_labels.font.color.rgb = WHITE
ch5.category_axis.tick_labels.font.size = Pt(13)
ch5.category_axis.tick_labels.font.bold = True
style_chart_text(ch5, color=WHITE)

add_rect(s, Inches(0.9), Inches(6.45), Inches(4.4), Inches(0.7),
         color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.3
add_text(s, Inches(1.1), Inches(6.45), Inches(4.1), Inches(0.7),
         [[("56,8% ", {'size': 15, 'bold': True, 'color': PURPLE}),
           ("przyznaje: marka wpływa na zakup", {'size': 13, 'color': INK})]],
         anchor=MSO_ANCHOR.MIDDLE)
page_number(s, 9, dark=True)


# ===========================================================================
# SLAJD 10 — WERYFIKACJA HIPOTEZ
# ===========================================================================
s = base_slide()
kicker(s, "Podsumowanie wyników", TEAL)
section_title(s, "Weryfikacja hipotez", TEAL)

hres = [
    ("H1", "iPhone = prestiż, Samsung = cena/bateria", "POTWIERDZONA", GREEN),
    ("H2", "iPhone nie jest symbolem statusu", "POTWIERDZONA", GREEN),
    ("H3", "Zmiana systemu mało istotna", "CZĘŚCIOWO", AMBER),
    ("H4", "Przy braku ograniczeń wybór: iPhone", "POTWIERDZONA", GREEN),
]
hy = Inches(2.6)
for tag, txt, status, col in hres:
    add_rect(s, Inches(0.9), hy, Inches(11.5), Inches(0.95),
             color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True).adjustments[0] = 0.16
    add_circle(s, Inches(1.08), hy + Inches(0.15), Inches(0.65), color=PURPLE)
    add_text(s, Inches(1.08), hy + Inches(0.15), Inches(0.65), Inches(0.65),
             [[(tag, {'size': 17, 'bold': True, 'color': WHITE})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.0), hy, Inches(7.0), Inches(0.95),
             [[(txt, {'size': 16, 'color': INK})]], anchor=MSO_ANCHOR.MIDDLE)
    badge = add_rect(s, Inches(9.3), hy + Inches(0.2), Inches(2.9), Inches(0.55),
                     color=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    badge.adjustments[0] = 0.5
    add_text(s, Inches(9.3), hy + Inches(0.2), Inches(2.9), Inches(0.55),
             [[(status, {'size': 14, 'bold': True, 'color': WHITE})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    hy += Inches(1.08)
page_number(s, 10)


# ===========================================================================
# SLAJD 11 — WNIOSKI
# ===========================================================================
s = base_slide()
kicker(s, "Rozdział 3 — Wnioski", CORAL)
section_title(s, "Najważniejsze wnioski", CORAL)

concl = [
    ("iPhone = prestiż", "Marka premium: jakość, design, aparat. Słabości: bateria i cena.", IPHONE_C),
    ("Samsung = praktyczność", "Funkcjonalność, bateria, intuicyjność i korzystna cena.", SAMSUNG_C),
    ("Status traci znaczenie", "Konsumenci coraz częściej oceniają realną użyteczność, nie prestiż.", TEAL),
    ("Marka napędza zakupy", "Marka to jeden z kluczowych czynników decyzji konsumenckich.", PURPLE),
]
gx, gy = Inches(0.9), Inches(2.55)
cw, ch = Inches(5.7), Inches(1.95)
for i, (title, txt, col) in enumerate(concl):
    x = gx + (cw + Inches(0.15)) * (i % 2)
    y = gy + (ch + Inches(0.2)) * (i // 2)
    card = add_rect(s, x, y, cw, ch, color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    card.adjustments[0] = 0.07
    add_circle(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55), color=col)
    add_text(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55), Inches(0.55),
             [[(str(i+1), {'size': 18, 'bold': True, 'color': WHITE})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(1.05), y + Inches(0.28), cw - Inches(1.3), Inches(0.6),
             [[(title, {'size': 18, 'bold': True, 'color': col})]])
    add_text(s, x + Inches(0.35), y + Inches(1.0), cw - Inches(0.65), Inches(0.85),
             [[(txt, {'size': 14, 'color': INK})]], line_spacing=1.05)
page_number(s, 11)


# ===========================================================================
# SLAJD 12 — ZALECENIA PRAKTYCZNE + PODZIĘKOWANIE
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, grad=(PURPLE, PURPLE_DK), grad_angle=120)
add_circle(s, SW - Inches(2.2), Inches(-1.0), Inches(3.4), grad=(AMBER, CORAL))
add_circle(s, Inches(-1.2), SH - Inches(2.0), Inches(3.6), grad=(TEAL, SKY))

kicker(s, "Rozdział 3 — Zalecenia", AMBER)
add_text(s, Inches(0.9), Inches(1.15), Inches(11.5), Inches(0.9),
         [[("Zalecenia praktyczne", {'size': 32, 'bold': True, 'color': WHITE})]])

recs = [
    ("Apple", "poprawić baterię oraz relację jakości do ceny.", PINK),
    ("Samsung", "podkreślać funkcjonalność, budować prestiż marki.", TEAL),
    ("Obie marki", "rozwijać lojalność i komunikację marketingową.", AMBER),
    ("Komunikacja", "akcentować konkretne, racjonalne korzyści użytkowe.", CORAL),
]
ry = Inches(2.3)
for tag, txt, col in recs:
    add_rect(s, Inches(0.9), ry, Inches(7.6), Inches(0.8),
             color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.25
    add_rect(s, Inches(0.9), ry, Inches(0.14), Inches(0.8), color=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(1.25), ry, Inches(7.1), Inches(0.8),
             [[(tag + ": ", {'size': 14.5, 'bold': True, 'color': col}),
               (txt, {'size': 14, 'color': INK})]], anchor=MSO_ANCHOR.MIDDLE)
    ry += Inches(0.95)

add_text(s, Inches(8.9), Inches(2.6), Inches(3.8), Inches(3.2),
         [[("Dziękujemy", {'size': 40, 'bold': True, 'color': WHITE})],
          [("za uwagę!", {'size': 30, 'bold': True, 'color': AMBER})],
          [("", {'size': 10})],
          [("Pytania?", {'size': 18, 'italic': True, 'color': RGBColor(0xE7, 0xDE, 0xFB)})]],
         space_after=2)

add_text(s, Inches(0.9), SH - Inches(0.7), Inches(11.5), Inches(0.4),
         [[("UMCS w Lublinie  •  Wydział Ekonomiczny  •  Zarządzanie  •  Badanie marketingowe 2026",
            {'size': 11, 'color': RGBColor(0xCB, 0xBD, 0xF0)})]])


# ---------------------------------------------------------------------------
prs.save('Prezentacja_iPhone_vs_Samsung.pptx')
print("OK - zapisano Prezentacja_iPhone_vs_Samsung.pptx  | slajdów:", len(prs.slides._sldIdLst))
